"""Listado de créditos aprobados que entregó la gobernación, cruzado en vivo con
la consulta real de cada expediente (estado de carpeta, monto del sistema,
técnico) y con el detalle del estado que se va cargando a mano.

El listado en sí (nro, proyecto, monto, línea, fecha de firma) es el de la nota
de la gobernación y no se edita; lo que se edita es `detalle_estado`, que vive en
la consulta y se guarda con el PATCH genérico de /api/consultas/{id}.
"""
import re
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import text

from db import engine
from auth import require_login
from formatos import _monto, _dmy
from constantes import ESTADOS_CARPETA, ESTADO_CARPETA_COLOR, SITUACIONES_TRAMITE, SITUACION_TRAMITE_COLOR

router = APIRouter(prefix="/api/listado-gobernacion", tags=["gobernacion"])


@router.get("")
def listar(usuario=Depends(require_login)):
    with engine.connect() as conn:
        rows = conn.execute(text("""
            SELECT g.nro, g.solicitante, g.proyecto, g.monto, g.linea, g.fecha_firma, g.nota, g.anulado, g.omitir_informe,
                   c.id AS consulta_id, c.codigo, c.tecnico, c.instancia_superior,
                   c.estado_carpeta, c.detalle_estado, c.situacion_tramite, c.firmado,
                   COALESCE(NULLIF(c.monto_instancia_superior, 0), NULLIF(c.monto_confirmado, 0), c.monto) AS monto_sistema
            FROM sde_listado_gobernacion g
            LEFT JOIN sde_consultas c ON c.id = g.consulta_id
            ORDER BY g.nro
        """)).mappings().all()

    casos = []
    for r in rows:
        sistema = int(r["monto_sistema"] or 0) if r["consulta_id"] else None
        casos.append({
            "nro": r["nro"], "solicitante": r["solicitante"], "proyecto": r["proyecto"] or "",
            "monto": r["monto"], "monto_fmt": _monto(r["monto"]),
            "linea": r["linea"] or "—",
            "fecha_firma": _dmy(r["fecha_firma"]) if r["fecha_firma"] else "—",
            "nota": r["nota"] or "",
            "anulado": bool(r["anulado"]),
            "omitir_informe": bool(r["omitir_informe"]),
            "consulta_id": r["consulta_id"], "codigo": r["codigo"],
            "tecnico": r["tecnico"] or "Sin asignar" if r["consulta_id"] else "—",
            "en_instancia_superior": bool(r["instancia_superior"]),
            "estado_carpeta": r["estado_carpeta"] or "",
            "detalle_estado": r["detalle_estado"] or "",
            "situacion_tramite": r["situacion_tramite"] or "",
            "monto_sistema": sistema,
            "monto_sistema_fmt": _monto(sistema) if sistema is not None else "—",
            "monto_difiere": sistema is not None and r["monto"] is not None and sistema != r["monto"],
        })
    return {
        "total": len(casos),
        "monto_total_fmt": _monto(sum(c["monto"] or 0 for c in casos)),
        "estados": ESTADOS_CARPETA,
        "estado_carpeta_color": ESTADO_CARPETA_COLOR,
        "situaciones": SITUACIONES_TRAMITE,
        "situacion_color": SITUACION_TRAMITE_COLOR,
        "casos": casos,
    }


# Orden y rótulos del Excel — los mismos criterios que el informe imprimible
# (de lo más avanzado a lo que recién empieza; estados de salida al final).
_ORDEN = ["DESEMBOLSADO", "EN TRAMITE DE DESEMBOLSO", "EN TRAMITE EN CFI SEDE",
          "FIRMADO - EN TRAMITE PARA INGRESAR A CFI (DOCUMENTACION O SGR)",
          "ENVIADO A FIRMA DE REPRESENTANTE", "DESISTIDO POR SOLICITANTE",
          "RECHAZADO POR ANALISIS DE RIESGO", ""]
_ROTULO = {"DESEMBOLSADO": "Desembolsado", "EN TRAMITE DE DESEMBOLSO": "En trámite de desembolso",
           "EN TRAMITE EN CFI SEDE": "En trámite en CFI Sede",
           "FIRMADO - EN TRAMITE PARA INGRESAR A CFI (DOCUMENTACION O SGR)": "Firmado - en trámite para ingresar a CFI",
           "ENVIADO A FIRMA DE REPRESENTANTE": "Enviado a firma del representante",
           "DESISTIDO POR SOLICITANTE": "Desistido por el solicitante",
           "RECHAZADO POR ANALISIS DE RIESGO": "Rechazado por análisis de riesgo", "": "Sin clasificar"}
_FONDO = {"DESEMBOLSADO": "E3F5EA", "EN TRAMITE DE DESEMBOLSO": "EAF7F0", "EN TRAMITE EN CFI SEDE": "EFE9FB",
          "FIRMADO - EN TRAMITE PARA INGRESAR A CFI (DOCUMENTACION O SGR)": "FBF3DC",
          "ENVIADO A FIRMA DE REPRESENTANTE": "FBE7EF"}


@router.get("/excel")
def exportar_excel(usuario=Depends(require_login)):
    """Excel para entregar a la gobernación: mismos expedientes que el informe
    imprimible (sin los reemplazados ni los omitidos), con el estado, su detalle
    y los datos de contacto para llamar al solicitante. Sin técnicos."""
    with engine.connect() as conn:
        rows = conn.execute(text("""
            SELECT g.nro, g.solicitante, g.monto, c.estado_carpeta, c.detalle_estado,
                   COALESCE(c.telefono, '') AS telefono, COALESCE(c.mail, '') AS mail
            FROM sde_listado_gobernacion g
            LEFT JOIN sde_consultas c ON c.id = g.consulta_id
            WHERE NOT COALESCE(g.anulado, FALSE) AND NOT COALESCE(g.omitir_informe, FALSE)
        """)).mappings().all()
    rows = sorted(rows, key=lambda r: (_ORDEN.index(r["estado_carpeta"] or "") if (r["estado_carpeta"] or "") in _ORDEN else len(_ORDEN), r["nro"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Créditos aprobados"
    fuente = "Arial"
    encabezado = ["N°", "Solicitante", "Teléfono", "Mail", "Monto", "Estado", "Detalle del estado"]
    ws.append(encabezado)
    for i in range(1, len(encabezado) + 1):
        c = ws.cell(1, i)
        c.font = Font(name=fuente, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1C2443")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    borde = Border(*(Side(style="thin", color="D9DCE3"),) * 4)
    for n, r in enumerate(rows, start=2):
        est = r["estado_carpeta"] or ""
        detalle = (r["detalle_estado"] or "").strip()
        if re.match(r"^desembolsad[oa]s?\.?$", detalle, re.I):
            detalle = ""  # repite el estado: no aporta
        ws.append([r["nro"], r["solicitante"], (r["telefono"] or "").strip(), (r["mail"] or "").strip(),
                   r["monto"], _ROTULO.get(est, est), detalle])
        for i in range(1, 8):
            x = ws.cell(n, i)
            x.font = Font(name=fuente, size=10)
            x.border = borde
            x.alignment = Alignment(vertical="top", wrap_text=i in (2, 4, 6, 7))
        if est in _FONDO:
            ws.cell(n, 6).fill = PatternFill("solid", fgColor=_FONDO[est])
        ws.cell(n, 3).number_format = "@"
        ws.cell(n, 5).number_format = "$#,##0"
    ultima = len(rows) + 1
    total = ultima + 1
    ws.cell(total, 1, f"=COUNTA(B2:B{ultima})")
    ws.cell(total, 2, "TOTAL")
    ws.cell(total, 5, f"=SUM(E2:E{ultima})")
    ws.cell(total, 5).number_format = "$#,##0"
    for i in range(1, 8):
        x = ws.cell(total, i)
        x.font = Font(name=fuente, bold=True, size=10)
        x.fill = PatternFill("solid", fgColor="EEF1F7")
        x.border = borde
    for col, ancho in zip("ABCDEFG", [6, 34, 20, 36, 16, 34, 60]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ultima}"

    buf = BytesIO()
    wb.save(buf)
    hoy = datetime.now(ZoneInfo("America/Argentina/Buenos_Aires")).strftime("%d-%m-%Y")
    nombre = f"Creditos aprobados - estado y contactos ({hoy}).xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )
